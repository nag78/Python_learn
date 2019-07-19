from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


wb = Workbook()
ws = wb.active
ws['A1'] = 'Серия 1'
for i in range(1, 11):
    ws.append([i])
chart = BarChart()
chart.title = "Первая серия данных."
data = Reference(ws, min_col=1, min_row=1,
                 max_col=1, max_row=11)

chart.add_data(data, titles_from_data=True)
ws.add_chart(chart, 'C2')
wb.save('sampleChart1.xlsx')
